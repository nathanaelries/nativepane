"""Large-file regressions using independent OOXML writers/readers and a browser.
QA-only dependencies: python-docx, openpyxl, Playwright and Chromium/Edge.
"""
import os
import time
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright, expect

root = Path(__file__).resolve().parent.parent
out = root / "tmp" / "qa"
out.mkdir(parents=True, exist_ok=True)
base = os.environ.get("NATIVEPANE_URL", "http://127.0.0.1:8080")

workbook = Workbook()
sheet = workbook.active
for row in range(1, 5001):
    sheet.append([row * 10 + col for col in range(10)])
xlsx = out / "large-original.xlsx"
workbook.save(xlsx)

document = Document()
document.add_heading("Large document regression", 0)
for paragraph in range(500):
    p = document.add_paragraph()
    for run in range(50):
        p.add_run(f"{run} ")
docx = out / "large-original.docx"
document.save(docx)

with sync_playwright() as p:
    options = {"headless": True}
    if os.environ.get("BROWSER_EXECUTABLE"):
        options["executable_path"] = os.environ["BROWSER_EXECUTABLE"]
    browser = p.chromium.launch(**options)
    page = browser.new_page(viewport={"width": 1440, "height": 1050})
    page.set_default_timeout(120000)
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.goto(base)

    for source, count in ((xlsx, 50000), (docx, 25001)):
        started = time.monotonic()
        page.locator("#upload").set_input_files(str(source))
        expect(page.locator("#filename")).to_have_text(source.name)
        expect(page.locator("#document-info")).to_contain_text(f"{count:,} fields")
        assert page.locator("#open-status").get_attribute("hidden") is not None
        if source.suffix == ".xlsx":
            # Paging bounds the live DOM even though all 50,000 cells are editable.
            assert page.locator("[data-field]").count() == 1000
            page.get_by_role("button", name="Rows →", exact=True).click()
            field = page.get_by_role("textbox", name="A101", exact=True)
        else:
            assert page.locator("[data-field]").count() == count
            field = page.locator("[data-field]").last
        field.fill("Large file edited")
        expect(page.locator("#save-status")).to_have_text("Saved · revision 2")
        with page.expect_download() as download:
            page.locator("#download").click()
        saved = out / ("large-saved" + source.suffix)
        download.value.save_as(str(saved))
        if source.suffix == ".xlsx":
            reopened = load_workbook(saved)
            assert reopened.active["A101"].value == "Large file edited"
            assert reopened.active["J5000"].value == 50009
            changed_part = "xl/worksheets/sheet1.xml"
        else:
            reopened = Document(saved)
            assert reopened.paragraphs[-1].runs[-1].text == "Large file edited"
            assert len(reopened.paragraphs) == 501
            changed_part = "word/document.xml"
        with ZipFile(source) as before, ZipFile(saved) as after:
            assert before.namelist() == after.namelist()
            for name in before.namelist():
                if name != changed_part:
                    assert before.read(name) == after.read(name), name
        page.locator("#close").click()
        expect(page.locator("#welcome")).to_be_visible()
        # The file input is reset so users can retry the very same file.
        page.locator("#upload").set_input_files(str(source))
        expect(page.locator("#workspace")).to_be_visible()
        page.locator("#close").click()
        print(f"{source.suffix.upper()}: {count:,} fields, edit/autosave/download/independent reopen/retry PASS ({time.monotonic()-started:.1f}s)")
    assert not errors, errors
    browser.close()
