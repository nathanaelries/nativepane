"""Optional QA: author and reopen using independent permissively licensed readers.
Requires python-docx, openpyxl, python-pptx in the QA environment, not the product.
"""
from io import BytesIO
from pathlib import Path
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches
from container_smoke import edit_file

out_dir = Path(__file__).resolve().parent.parent / "tmp" / "qa"
out_dir.mkdir(parents=True, exist_ok=True)
expected = "Verified edit & <round trip>"

doc = Document()
doc.add_heading("Independent Word document", 0)
p = doc.add_paragraph("First run. ")
p.add_run("Bold text retained.").bold = True
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Table content"
doc.sections[0].header.paragraphs[0].text = "Header must be preserved"
buf = BytesIO()
doc.save(buf)
source = buf.getvalue()
(out_dir / "independent.docx").write_bytes(source)
saved = edit_file("independent.docx", source)
reopened = Document(BytesIO(saved))
assert reopened.paragraphs[0].text == expected
assert reopened.paragraphs[1].runs[1].bold
assert reopened.tables[0].cell(0, 0).text == "Table content"
assert reopened.sections[0].header.paragraphs[0].text == "Header must be preserved"
(out_dir / "saved.docx").write_bytes(saved)
print("python-docx: independently authored and reopened; text, bold, table, header PASS")

workbook = Workbook()
sheet = workbook.active
sheet.title = "Original worksheet"
sheet["A1"] = "Edit this"
sheet["B2"] = 42
sheet["C2"] = "=B2*2"
sheet["B2"].number_format = "0.00"
workbook.create_sheet("Second worksheet")["A1"] = "Keep this"
buf = BytesIO()
workbook.save(buf)
source = buf.getvalue()
(out_dir / "independent.xlsx").write_bytes(source)
saved = edit_file("independent.xlsx", source)
reopened = load_workbook(BytesIO(saved))
assert reopened.active["A1"].value == expected
assert reopened.active["C2"].value == "=B2*2"
assert reopened.active["B2"].number_format == "0.00"
assert reopened["Second worksheet"]["A1"].value == "Keep this"
(out_dir / "saved.xlsx").write_bytes(saved)
print("openpyxl: independently authored and reopened; text, formula, style, second sheet PASS")

deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[6])
shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
shape.text = "Edit this slide"
slide = deck.slides.add_slide(deck.slide_layouts[6])
slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text = "Keep second slide"
buf = BytesIO()
deck.save(buf)
source = buf.getvalue()
(out_dir / "independent.pptx").write_bytes(source)
saved = edit_file("independent.pptx", source)
reopened = Presentation(BytesIO(saved))
assert reopened.slides[0].shapes[0].text == expected
assert reopened.slides[1].shapes[0].text == "Keep second slide"
assert reopened.slides[0].shapes[0].left == Inches(1)
(out_dir / "saved.pptx").write_bytes(saved)
print("python-pptx: independently authored and reopened; text, second slide, coordinates PASS")
