"""Focused original fixtures supplement the unmodified Office-saved corpus.
Unknown extensions are injected only while authoring test inputs; the product
must preserve their raw bytes. All generated inputs/downloads live under tmp/.
"""
import datetime
import io
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as E

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_WINDOWS_1900, CALENDAR_MAC_1904
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image

from container_smoke import request
from corpus_helpers import signature, verify_edit

out = Path(__file__).resolve().parent.parent / "tmp" / "preservation"
out.mkdir(parents=True, exist_ok=True)
S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def run_edits(name, source, select):
    access = json.loads(request("/api/v1/sessions", "POST", {"filename": name, "mode": "edit"}))
    route = "/api/v1/sessions/" + access["id"]
    try:
        request(route + "/file", "PUT", source)
        opened = json.loads(request(route + "/document", token=access["token"]))
        edits = select(opened["model"])
        request(route + "/document", "PATCH", {"revision": 1, "edits": edits}, access["token"])
        return request(route + "/file", token=access["token"])
    finally:
        request(route, "DELETE")


for epoch in (CALENDAR_WINDOWS_1900, CALENDAR_MAC_1904):
    workbook = Workbook()
    workbook.epoch = epoch
    sheet = workbook.active
    sheet["A1"] = datetime.datetime(2026, 9, 30, 14, 30)
    sheet["A1"].number_format = 'yyyy-mm-dd hh:mm'
    sheet["A2"], sheet["A2"].number_format = 1234.56, '$#,##0.00'
    sheet["A3"], sheet["A3"].number_format = .125, '0.0%'
    sheet["B1"] = "Merged label"
    sheet.merge_cells("B1:C2")
    sheet["D1"] = '=SUM(A2:A3)'
    buffer = io.BytesIO()
    workbook.save(buffer)
    extension = b'<extLst><ext uri="urn:nativepane:preservation"><np:payload xmlns:np="urn:nativepane:vendor" flag="preserve"> untouched &amp; data </np:payload></ext></extLst>'
    source_buffer = io.BytesIO()
    with ZipFile(buffer) as original, ZipFile(source_buffer, 'w', ZIP_DEFLATED) as authored:
        for member in original.infolist():
            raw = original.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                # A1 is the first cell, authored by openpyxl with a numeric value.
                raw = raw.replace(b'</c>', extension + b'</c>', 1)
            authored.writestr(member, raw)
    source = source_buffer.getvalue()
    name = f"formats-{epoch.year}.xlsx"
    (out / name).write_bytes(source)

    def select(model):
        fields = {f['address']: f for f in model['blocks'][0]['fields']}
        assert fields['D1']['readOnly']
        return [{'id': fields['A1']['id'], 'text': str(float(fields['A1']['text']) + 1)},
                {'id': fields['A2']['id'], 'text': '2345.67'},
                {'id': fields['A3']['id'], 'text': '0.5'},
                {'id': fields['B1']['id'], 'text': 'Changed merged label'}]

    saved = run_edits(name, source, select)
    with ZipFile(io.BytesIO(source)) as before, ZipFile(io.BytesIO(saved)) as after:
        assert before.namelist() == after.namelist()
        for member in before.namelist():
            if member != "xl/worksheets/sheet1.xml":
                assert before.read(member) == after.read(member)
        a, b = before.read("xl/worksheets/sheet1.xml"), after.read("xl/worksheets/sheet1.xml")
        assert extension in b, "unknown edited-cell extension lost"
        a, b = E.fromstring(a), E.fromstring(b)
        for root in (a, b):
            for cell in root.iter('{' + S + '}c'):
                if cell.get('r') in ('A1', 'A2', 'A3', 'B1'):
                    cell.attrib.pop('t', None)
                    for child in list(cell):
                        if child.tag in ('{' + S + '}v', '{' + S + '}is'):
                            cell.remove(child)
        assert signature(a) == signature(b), "cell properties, extension tree or merges changed"
    reopened = load_workbook(io.BytesIO(saved))
    s = reopened.active
    assert reopened.epoch == epoch
    assert s['A1'].value == datetime.datetime(2026, 10, 1, 14, 30)
    assert s['A1'].number_format == 'yyyy-mm-dd hh:mm'
    assert s['A2'].value == 2345.67 and s['A2'].number_format == '$#,##0.00'
    assert s['A3'].value == .5 and s['A3'].number_format == '0.0%'
    assert s['B1'].value == 'Changed merged label' and str(next(iter(s.merged_cells.ranges))) == 'B1:C2'
    assert s['D1'].value == '=SUM(A2:A3)'
    print(name + ': date epoch, custom/currency/percentage formats, merge, formula, edited-cell extension PASS')

presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[6])
shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
paragraph = shape.text_frame.paragraphs[0]
first = paragraph.add_run()
first.text = 'Bold colored run'
first.font.name, first.font.size, first.font.bold = 'Arial', Pt(32), True
first.font.color.rgb = RGBColor(0x24, 0x43, 0x5B)
second = paragraph.add_run()
second.text, second.font.size, second.font.italic = ' Italic smaller run', Pt(18), True
image = out / 'original-image.png'
Image.new('RGB', (16, 16), (20, 80, 140)).save(image)
slide.shapes.add_picture(str(image), Inches(1), Inches(3), Inches(1), Inches(1))
buffer = io.BytesIO()
presentation.save(buffer)
source = buffer.getvalue()
(out / 'rich-slide.pptx').write_bytes(source)
saved = run_edits('rich-slide.pptx', source, lambda model: [{'id': model['blocks'][0]['fields'][0]['id'], 'text': 'Edited rich run'}])
with ZipFile(io.BytesIO(source)) as before, ZipFile(io.BytesIO(saved)) as after:
    for member in before.namelist():
        if member != 'ppt/slides/slide1.xml':
            assert before.read(member) == after.read(member), member
    verify_edit(E.fromstring(before.read('ppt/slides/slide1.xml')), E.fromstring(after.read('ppt/slides/slide1.xml')), {'text': first.text}, 'Edited rich run', 'pptx')
reopened = Presentation(io.BytesIO(saved))
original_shape, changed_shape = presentation.slides[0].shapes[0], reopened.slides[0].shapes[0]
assert (original_shape.left, original_shape.top, original_shape.width, original_shape.height) == (changed_shape.left, changed_shape.top, changed_shape.width, changed_shape.height)
runs = changed_shape.text_frame.paragraphs[0].runs
assert runs[0].text == 'Edited rich run' and runs[0].font.bold and runs[0].font.size == Pt(32)
assert runs[0].font.color.rgb == RGBColor(0x24, 0x43, 0x5B) and runs[0].font.name == 'Arial'
assert runs[1].font.italic and runs[1].font.size == Pt(18) and runs[1].text == second.text
assert presentation.slides[0].shapes[1].image.blob == reopened.slides[0].shapes[1].image.blob
print('PPTX: rich text, shapes, coordinates, images, themes and package parts PASS')
