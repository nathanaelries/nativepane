"""DOCX table regression: independent authoring, browser layout, native edit/save.
QA-only dependencies: python-docx and Playwright. Outputs stay under tmp/qa.
"""
import os
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from playwright.sync_api import sync_playwright, expect

root = Path(__file__).resolve().parent.parent
out = root / "tmp" / "qa"
out.mkdir(parents=True, exist_ok=True)
base = os.environ.get("NATIVEPANE_URL", "http://127.0.0.1:8080")

document = Document()
document.add_paragraph("Before the table")
table = document.add_table(rows=4, cols=3)
table.style = "Table Grid"
table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"
table.cell(0, 2).text = "Third column"
table.cell(1, 0).merge(table.cell(2, 0)).text = "Vertical cell"
cell = table.cell(1, 1)
cell.text = "Editable cell"
cell.add_paragraph("Second paragraph")
cell.add_table(rows=1, cols=1).cell(0, 0).text = "Nested cell"
table.cell(2, 1).text = "Final row"
table.cell(2, 2).text = "Last cell"
table.cell(3, 0).text = "Full row"
table.cell(3, 1).text = "Middle"
table.cell(3, 2).text = "Right"
document.add_paragraph("After the table")
source = out / "independent-tables.docx"
document.save(source)

with sync_playwright() as p:
    options = {"headless": True}
    if os.environ.get("BROWSER_EXECUTABLE"):
        options["executable_path"] = os.environ["BROWSER_EXECUTABLE"]
    browser = p.chromium.launch(**options)
    page = browser.new_page(viewport={"width": 1440, "height": 1100})
    errors = []
    page.on("pageerror", lambda e: errors.append(str(e)))
    page.goto(base)
    page.locator("#upload").set_input_files(str(source))
    expect(page.locator("#filename")).to_have_text(source.name)
    expect(page.locator("table.document-table")).to_have_count(2)
    outer = page.locator(".page-content > table.document-table")
    expect(outer.locator(":scope > tbody > tr")).to_have_count(4)
    expect(outer.locator(":scope > tbody > tr").nth(0).locator(":scope > td").first).to_have_attribute("colspan", "2")
    expect(outer.locator(":scope > tbody > tr").nth(1).locator(":scope > td").first).to_have_attribute("rowspan", "2")
    expect(outer.locator(":scope > tbody > tr").nth(2).locator(":scope > td")).to_have_count(2)
    # The nested cell appears once in the DOM and stays inside the parent cell.
    expect(page.locator('[data-field]').filter(has_text="Nested cell")).to_have_count(1)
    owner = outer.locator(":scope > tbody > tr").nth(1).locator(":scope > td").nth(1)
    expect(owner.locator("table.document-table")).to_have_count(1)
    expect(owner.locator(":scope > p").nth(1)).to_have_text("Second paragraph")
    flow = page.locator(".page-content").first
    expect(flow.locator(":scope > p").first).to_have_text("Before the table")
    expect(flow.locator(":scope > p").last).to_have_text("After the table")
    # Coordinate checks catch visually flattened/shifted rows even if text survives.
    header = outer.locator(":scope > tbody > tr").nth(0).locator(":scope > td").first.bounding_box()
    third = outer.locator(":scope > tbody > tr").nth(0).locator(":scope > td").nth(1).bounding_box()
    assert header["width"] > third["width"] * 1.7
    assert abs(header["y"] - third["y"]) < 1
    assert third["x"] >= header["x"] + header["width"] - 1
    vertical = outer.locator(":scope > tbody > tr").nth(1).locator(":scope > td").first.bounding_box()
    final = outer.locator(":scope > tbody > tr").nth(2).locator(":scope > td").first.bounding_box()
    assert final["x"] >= vertical["x"] + vertical["width"] - 1
    assert vertical["y"] + vertical["height"] >= final["y"] + final["height"] - 1
    page.locator('[data-field]').filter(has_text="Editable cell").fill("Updated table cell")
    page.locator('[data-field]').filter(has_text="Nested cell").fill("Updated nested cell")
    expect(page.locator("#save-status")).to_have_text("Saved · revision 2", timeout=15000)
    page.get_by_role("button", name="↶ Undo", exact=True).click()
    expect(page.locator('[data-field]').filter(has_text="Nested cell")).to_have_count(1)
    page.get_by_role("button", name="↷ Redo", exact=True).click()
    expect(page.locator('[data-field]').filter(has_text="Updated nested cell")).to_have_count(1)
    with page.expect_download() as download:
        page.get_by_role("button", name="↓ Download original format", exact=True).click()
    destination = out / "browser-tables-saved.docx"
    download.value.save_as(str(destination))
    page.screenshot(path=str(out / "docx-tables.png"), full_page=True)
    reopened = Document(destination)
    assert reopened.tables[0].cell(1, 1).paragraphs[0].text == "Updated table cell"
    assert reopened.tables[0].cell(1, 1).tables[0].cell(0, 0).text == "Updated nested cell"
    assert reopened.tables[0].cell(0, 0)._tc is reopened.tables[0].cell(0, 1)._tc
    assert reopened.tables[0].cell(1, 0)._tc is reopened.tables[0].cell(2, 0)._tc
    with ZipFile(source) as before, ZipFile(destination) as after:
        assert set(before.namelist()) == set(after.namelist())
        assert [n for n in before.namelist() if before.read(n) != after.read(n)] == ["word/document.xml"]
    assert not errors, errors
    page.get_by_role("button", name="Close pane", exact=True).click()
    browser.close()
    print("DOCX tables: rows, nested tables, merges, flow order, layout, undo/redo, native save and independent reopen PASS")
