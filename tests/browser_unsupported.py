"""Original supplemental DOCX: omitted constructs must remain visibly marked."""
import os
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image
from playwright.sync_api import sync_playwright, expect

root = Path(__file__).resolve().parent.parent
out = root / 'tmp' / 'qa'
out.mkdir(parents=True, exist_ok=True)
image = out / 'marker-image.png'
Image.new('RGB', (16, 16), (20, 80, 140)).save(image)
doc = Document()
doc.add_paragraph('Visible text before the drawing').add_run().add_picture(str(image), width=Inches(.5))
doc.add_paragraph('Inherited list numbering', style='List Bullet')
p = doc.add_paragraph('Text with a tab and an inline break')
p.add_run().add_tab()
p.add_run().add_break()
math = OxmlElement('m:oMath')
math_run = OxmlElement('m:r')
math_text = OxmlElement('m:t')
math_text.text = 'x + y'
math_run.append(math_text)
math.append(math_run)
p._p.append(math)
table = doc.add_table(rows=1, cols=1)
table.cell(0, 0).text = 'Floating table content'
position = OxmlElement('w:tblpPr')
for key, value in {'horzAnchor':'margin', 'vertAnchor':'text', 'tblpX':'0', 'tblpY':'0'}.items():
    position.set(qn('w:' + key), value)
table._tbl.tblPr.append(position)
doc.sections[0].header.paragraphs[0].text = 'Header content retained in the source'
source = out / 'unsupported-regression.docx'
doc.save(source)

with sync_playwright() as p:
    options = {'headless': True}
    if os.environ.get('BROWSER_EXECUTABLE'):
        options['executable_path'] = os.environ['BROWSER_EXECUTABLE']
    browser = p.chromium.launch(**options)
    page = browser.new_page(viewport={'width':1440, 'height':1100}, locale='en-US', timezone_id='UTC')
    page.goto(os.environ.get('NATIVEPANE_URL', 'http://127.0.0.1:8080'))
    page.locator('#upload').set_input_files(str(source))
    expect(page.locator('#filename')).to_have_text(source.name)
    for marker in ('Image, drawing or embedded object not rendered', 'List numbering not rendered', 'Tab positioning not rendered', 'Inline line/page break not rendered', 'Equation not rendered', 'Floating table positioning not rendered', 'Header content not rendered'):
        expect(page.locator('.unsupported-marker').filter(has_text=marker)).to_be_visible()
    expect(page.locator('[data-field]').filter(has_text='Floating table content')).to_have_count(1)
    expect(page.locator('#save')).to_be_enabled()
    assert page.locator('[data-field]').filter(has_text='Header content retained').count() == 0
    page.evaluate('async () => {await document.fonts.ready; document.activeElement.blur();}')
    page.locator('#workspace').screenshot(path=str(out / 'unsupported-regression.png'), animations='disabled', caret='hide')
    browser.close()
    print('Unsupported Word constructs: inherited numbering, image, tabs/breaks, equation, floating table, header markers PASS')
