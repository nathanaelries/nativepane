"""Independent DOCX typography and measured pagination regression tests.
QA-only dependencies: python-docx and Playwright; no application dependencies.
"""
import os
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from playwright.sync_api import sync_playwright, expect

root = Path(__file__).resolve().parent.parent
out = root / "tmp" / "qa"
out.mkdir(parents=True, exist_ok=True)
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.75)
section.left_margin = section.right_margin = Inches(.75)
normal = document.styles["Normal"]
normal.font.name, normal.font.size = "Calibri", Pt(9.5)
normal.paragraph_format.line_spacing = 1.1
normal.paragraph_format.space_after = Pt(6)
title_style = document.styles["Title"]
title_style.font.name, title_style.font.size, title_style.font.bold = "Calibri", Pt(18), True
title_style.paragraph_format.line_spacing = 1
title_style.paragraph_format.space_after = Pt(10)
border = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
for key, value in {"val":"single", "sz":"8", "color":"4F81BD", "space":"4"}.items():
    bottom.set(qn("w:" + key), value)
border.append(bottom)
title_style.element.get_or_add_pPr().append(border)
heading_style = document.styles["Heading 1"]
heading_style.font.name, heading_style.font.size = "Calibri", Pt(12)
heading_style.paragraph_format.keep_with_next = True
document.add_paragraph("NativePane layout regression", style="Title")
document.add_paragraph("This original test document checks font inheritance, compact cells, page metrics and table continuation.")
document.add_paragraph("Long table", style="Heading 1")
table = document.add_table(rows=37, cols=3)
table.style = "Table Grid"
header = table.rows[0]._tr.get_or_add_trPr()
header.append(OxmlElement("w:tblHeader"))
for index, cell in enumerate(table.rows[0].cells):
    cell.text = ("Record", "Current status", "Notes")[index]
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "24435B")
    cell._tc.get_or_add_tcPr().append(shading)
    for run in cell.paragraphs[0].runs:
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("FFFFFF")
for i in range(1, 37):
    table.cell(i, 0).text = f"Item {i:02}"
    table.cell(i, 1).text = "Pending review. This deliberately long sentence verifies that cell wrapping uses the document font and the available column width."
    table.cell(i, 2).text = "Discuss the proposed change and record the next action. Preserve this note when editing another cell."
for row in table.rows:
    for col, cell in enumerate(row.cells):
        cell.width = Inches((1.4, 2.8, 2.8)[col])
        for para in cell.paragraphs:
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing = 1.08
table.cell(3, 0).merge(table.cell(4, 0))
document.add_paragraph("Following section", style="Heading 1")
document.add_paragraph("The paragraph after the table must remain visible, with no lost or duplicated editable fields.")
source = out / "layout-regression.docx"
document.save(source)

with sync_playwright() as p:
    options = {"headless": True}
    if os.environ.get("BROWSER_EXECUTABLE"):
        options["executable_path"] = os.environ["BROWSER_EXECUTABLE"]
    browser = p.chromium.launch(**options)
    page = browser.new_page(viewport={"width":1100,"height":1000})
    errors = []
    page.on("pageerror", lambda e: errors.append(str(e)))
    page.goto(os.environ.get("NATIVEPANE_URL", "http://127.0.0.1:8080"))
    page.locator("#upload").set_input_files(str(source))
    expect(page.locator("#filename")).to_have_text(source.name)
    pages = page.locator(".page")
    assert pages.count() >= 2
    assert page.locator(".page-content").first.locator("table").count() == 1, "Table left page one blank"
    assert page.locator('[data-repeated="true"]').count() >= 1, "Table header not repeated"
    metrics = page.evaluate("""() => {
      const title = document.querySelector('.page-content > p');
      const fields = [...document.querySelectorAll('[data-field]')];
      const wrap = document.querySelector('.canvas-wrap');
      return { titleSize:parseFloat(getComputedStyle(title).fontSize), titleWeight:getComputedStyle(title).fontWeight,
        titleBorder:getComputedStyle(title).borderBottomColor, paperWidth:parseFloat(document.querySelector('.page').style.width),
        bodySize:parseFloat(getComputedStyle(fields.find(x=>x.textContent==='Item 01')).fontSize),
        fieldCount:fields.length, uniqueFields:new Set(fields.map(x=>x.dataset.field)).size,
        overflow:[...document.querySelectorAll('.page-content')].map(x=>x.scrollHeight-((1056-144)-2)),
        repeatedEditable:[...document.querySelectorAll('[data-repeated] [contenteditable=true]')].length,
        canvasWidth:wrap.clientWidth, canvasScrollWidth:wrap.scrollWidth };
    }""")
    assert abs(metrics["titleSize"] - 24) < .1 and metrics["titleWeight"] == "700", metrics
    assert metrics["titleBorder"] == "rgb(79, 129, 189)" and metrics["paperWidth"] == 816, metrics
    assert abs(metrics["bodySize"] - 12.6667) < .1, metrics
    assert metrics["fieldCount"] == metrics["uniqueFields"], metrics
    assert max(metrics["overflow"]) <= 1, metrics
    assert metrics["repeatedEditable"] == 0, metrics
    assert metrics["canvasScrollWidth"] <= metrics["canvasWidth"] + 2, metrics
    header_row = page.locator(".page-content > table").first.locator(":scope > tbody > tr").first
    a, b = header_row.locator(":scope > td").nth(0).bounding_box(), header_row.locator(":scope > td").nth(1).bounding_box()
    assert 1.85 < b["width"] / a["width"] < 2.15, "Preferred cell widths ignored"
    expected_fields = metrics["uniqueFields"]
    before_pages = pages.count()
    # Enlarge a cell beyond one row's original height: pagination must update
    # during editing while keeping the caret on that cell and its ID unchanged.
    field = page.locator('[data-field]').filter(has_text="Item 20")
    field_id = field.get_attribute("data-field")
    replacement = "Expanded cell for layout reflow. " * 22
    field.fill(replacement)
    expect(page.locator("#save-status")).to_have_text("Saved · revision 2", timeout=15000)
    assert page.evaluate("document.activeElement.dataset.field") == field_id
    assert page.locator(".page").count() >= before_pages
    assert page.locator('[data-field]').count() == expected_fields
    with page.expect_download() as download:
        page.get_by_role("button", name="↓ Download original format", exact=True).click()
    destination = out / "layout-regression-saved.docx"
    download.value.save_as(str(destination))
    reopened = Document(destination)
    assert reopened.tables[0].cell(20, 0).text == replacement
    with ZipFile(source) as before, ZipFile(destination) as after:
        for part in ("word/styles.xml", "word/theme/theme1.xml"):
            assert before.read(part) == after.read(part)
    page.screenshot(path=str(out / "layout-regression.png"), full_page=True)
    assert not errors, errors
    browser.close()
    print("DOCX layout: inherited typography, page metrics, cell widths, table continuation, repeated headers, unique fields, edit reflow, native save PASS")
