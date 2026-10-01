"""Independent semantic and preservation checks for the pinned Office corpus."""
import hashlib
import io
import json
import urllib.error
import xml.etree.ElementTree as E
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from pptx import Presentation
from container_smoke import request
from corpus_helpers import verify_edit

root = Path(__file__).resolve().parent.parent
fixtures = root / "tests" / "fixtures"
out = root / "tmp" / "corpus"
out.mkdir(parents=True, exist_ok=True)
manifest = json.loads((fixtures / "manifest.json").read_text())


def document_text(container):
    texts = [p.text for p in container.paragraphs]
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(document_text(cell))
    return texts


for fixture in manifest["fixtures"]:
    source = (fixtures / fixture["file"]).read_bytes()
    assert hashlib.sha256(source).hexdigest() == fixture["sha256"]
    access = json.loads(request("/api/v1/sessions", "POST", {"filename": Path(fixture["file"]).name, "mode": "edit"}))
    route = "/api/v1/sessions/" + access["id"]
    try:
        request(route + "/file", "PUT", source)
        opened = json.loads(request(route + "/document", token=access["token"]))
        model = opened["model"]
        expected = json.loads((fixtures / fixture["semantic"]).read_text())
        assert model == expected, fixture["id"] + ": model differs from reviewed semantic baseline"
        fields = [f for b in model["blocks"] for f in b["fields"]]
        assert len({f["id"] for f in fields}) == len(fields)
        if fixture["readOnly"]:
            assert opened["mode"] == "view" and all(f.get("readOnly") for f in fields)
            try:
                request(route + "/document", "PATCH", {"revision": 1, "edits": [{"id": fields[0]["id"], "text": "Refused edit"}]}, access["token"])
                raise AssertionError("restricted edit accepted")
            except urllib.error.HTTPError as error:
                assert error.code == 403
            assert request(route + "/file", token=access["token"]) == source
        elif any(not f.get("readOnly") for f in fields):
            field = next(f for f in fields if not f.get("readOnly"))
            text = "77.5" if field.get("kind") == "number" else "Corpus verified edit"
            request(route + "/document", "PATCH", {"revision": 1, "edits": [{"id": field["id"], "text": text}]}, access["token"])
            saved = request(route + "/file", token=access["token"])
            (out / Path(fixture["file"]).name).write_bytes(saved)
            with ZipFile(io.BytesIO(source)) as before, ZipFile(io.BytesIO(saved)) as after:
                assert before.namelist() == after.namelist()
                changed = [n for n in before.namelist() if before.read(n) != after.read(n)]
                assert len(changed) == 1
                verify_edit(E.fromstring(before.read(changed[0])), E.fromstring(after.read(changed[0])), field, text, fixture["format"])
            if fixture["format"] == "docx":
                reopened = Document(io.BytesIO(saved))
                assert text in "\n".join(document_text(reopened))
            elif fixture["format"] == "xlsx":
                original = load_workbook(io.BytesIO(source))
                reopened = load_workbook(io.BytesIO(saved))
                assert reopened.sheetnames == original.sheetnames
                sheet_index = next(i for i, b in enumerate(model['blocks']) if any(f['id'] == field['id'] for f in b['fields']))
                cell = reopened.worksheets[sheet_index][field['address']]
                expected_value = (from_excel(77.5, reopened.epoch) if cell.is_date else 77.5) if text == '77.5' else text
                assert cell.value == expected_value
                for a, b in zip(original.worksheets, reopened.worksheets):
                    assert {str(r) for r in a.merged_cells.ranges} == {str(r) for r in b.merged_cells.ranges}
                    for row in a:
                        for cell in row:
                            other = b[cell.coordinate]
                            assert cell.number_format == other.number_format and cell.style_id == other.style_id
            else:
                original = Presentation(io.BytesIO(source))
                reopened = Presentation(io.BytesIO(saved))
                assert len(original.slides) == len(reopened.slides)
                assert (original.slide_width, original.slide_height) == (reopened.slide_width, reopened.slide_height)
                texts = []
                for a, b in zip(original.slides, reopened.slides):
                    assert len(a.shapes) == len(b.shapes)
                    for x, y in zip(a.shapes, b.shapes):
                        assert (x.shape_id, x.left, x.top, x.width, x.height, x.shape_type) == (y.shape_id, y.left, y.top, y.width, y.height, y.shape_type)
                        if x.shape_type == 13:
                            assert x.image.blob == y.image.blob
                        if y.has_text_frame:
                            texts.append(y.text)
                assert any(text in t for t in texts)
        else:
            # Image/text-box-only and formula-only Office projections have no
            # supported edit target. Their complete model and original download
            # are still gated, rather than inventing an editable field.
            assert request(route + '/file', token=access['token']) == source
        print(fixture["id"] + ": semantics, edit policy, native save and preservation PASS")
    finally:
        request(route, "DELETE")
